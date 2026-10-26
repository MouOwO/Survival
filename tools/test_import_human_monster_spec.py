"""Exercise both source documents and reject mismatched versions before generation."""
import copy
import json
import tempfile
import unittest
from pathlib import Path

import openpyxl
import import_human_monster_spec as importer


class HumanMonsterImportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.json_path = importer.SOURCE / 'human_spec.json'
        cls.xlsx_path = importer.SOURCE / '接入规格_人话版_v3.xlsx'
        cls.data = importer.load(cls.json_path, cls.xlsx_path)

    def test_complete_mode_ranges_and_workbook_only_bases(self):
        self.assertEqual(sum(map(len, self.data['modes'].values())), 235)
        self.assertEqual({key:len(rows) for key,rows in self.data['modes'].items()},
                         dict(standard=10, pure=10, rich=30, cold=50, scourge=50, ruins=50, celestial=35))
        self.assertEqual(set(self.data['waves']), set(range(1,31)))
        self.assertEqual(self.data['waves'][8]['health'], 4000)
        self.assertEqual(self.data['special_rotation'][8][0]['health'], 1105)
        self.assertGreater(self.data['sources']['cross_checked_cells'], 2800)

    def test_workbook_changes_drive_wave_base_even_when_json_is_unchanged(self):
        with tempfile.TemporaryDirectory() as folder:
            destination = Path(folder) / 'workbook.xlsx'
            book = openpyxl.load_workbook(self.xlsx_path)
            book['波次小怪基础值'].cell(9, 3).value = 4123.5
            book.save(destination)
            book.close()
            result = importer.load(self.json_path, destination)
        self.assertEqual(result['waves'][8]['health'], 4123.5)
        self.assertEqual(result['special_rotation'][8][0]['health'], 1105)

    def test_different_json_version_is_rejected(self):
        raw = json.loads(self.json_path.read_text(encoding='utf-8-sig'))
        raw['modes'][0]['HP倍率'] = 999
        with tempfile.TemporaryDirectory() as folder:
            destination = Path(folder) / 'human.json'
            destination.write_text(json.dumps(raw), encoding='utf-8')
            with self.assertRaisesRegex(AssertionError, 'Conflicting sources'):
                importer.load(destination, self.xlsx_path)

    def test_workbook_requires_every_wave(self):
        with tempfile.TemporaryDirectory() as folder:
            destination = Path(folder) / 'workbook.xlsx'
            book = openpyxl.load_workbook(self.xlsx_path)
            book['波次小怪基础值'].delete_rows(31)
            book.save(destination)
            book.close()
            with self.assertRaisesRegex(AssertionError, 'thirty wave bases'):
                importer.load(self.json_path, destination)

    def test_corrections_are_already_in_the_imported_values(self):
        bosses = self.data['attack_bosses']
        self.assertEqual((bosses[5]['health'], bosses[5]['war3_armor']), (10005, 98))
        self.assertEqual(bosses[20]['attack'], 60600)
        self.assertEqual(bosses[30]['health'], 40000015)
        archive = {row['name']:row for row in self.data['save_challenges']}
        self.assertEqual(archive['虚空之影1']['attack'], 189002)
        self.assertEqual(archive['神兽狩猎-古尔丹']['attack'], 5500002)
        self.assertEqual(archive['神兽狩猎-黑龙王子']['health'], 10080000000)
        self.assertEqual(archive['神兽狩猎-黑龙王子']['attack'], 168000002)

    def test_confirmed_wave_eight_health_is_scoped_to_rich_n10(self):
        self.assertEqual(self.data['special_rotation_health_overrides'], [{
            'mode_id': 'rich', 'difficulty': 10, 'wave_number': 8, 'position': 2,
            'health': 1420000, 'evidence': '用户确认2026-10-26丰饶N10第8波飞行小怪142万',
        }])
        self.assertEqual(self.data['special_rotation'][8][0]['health'], 1105)
        self.assertEqual(self.data['modes']['rich'][10]['hp_multiplier'], 356)

    def test_health_override_documents_must_match(self):
        raw = json.loads(self.json_path.read_text(encoding='utf-8-sig'))
        raw['special_rotation_health_overrides'][0]['最终HP'] = 1420001
        with tempfile.TemporaryDirectory() as folder:
            destination = Path(folder) / 'human.json'
            destination.write_text(json.dumps(raw, ensure_ascii=False), encoding='utf-8')
            with self.assertRaisesRegex(AssertionError, 'Conflicting sources: 轮换怪血量修正:2:最终HP'):
                importer.load(destination, self.xlsx_path)

    def test_health_override_workbook_row_cannot_be_omitted(self):
        with tempfile.TemporaryDirectory() as folder:
            destination = Path(folder) / 'workbook.xlsx'
            book = openpyxl.load_workbook(self.xlsx_path)
            book['轮换怪血量修正'].delete_rows(2)
            book.save(destination)
            book.close()
            with self.assertRaisesRegex(AssertionError, 'Conflicting row counts: special_rotation_health_overrides'):
                importer.load(self.json_path, destination)

    def test_health_override_rejects_invalid_or_unused_targets(self):
        raw = json.loads(self.json_path.read_text(encoding='utf-8-sig'))
        row = raw['special_rotation_health_overrides'][0]
        rotation_modes = {
            (source['波次'], 3 if '第3只' in source['出现规则'] else 2):
                {importer.MODE_IDS[name] for name in source['适用模式'].split('/')}
            for source in raw['special_rotation']
        }
        cases = [
            ('模式', '不存在', 'mode'), ('模式', '天灾', 'rotation slot'),
            ('难度', 31, 'difficulty'), ('难度', 1.5, '难度'), ('难度', True, '难度'),
            ('波次', 31, 'wave'), ('波次', 7, 'rotation slot'),
            ('轮换位置', 3, 'rotation slot'), ('轮换位置', 0, '轮换位置'),
            ('最终HP', 0, 'health'), ('最终HP', -1, 'health'),
            ('最终HP', float('nan'), 'health'), ('最终HP', float('inf'), 'health'),
            ('最终HP', True, 'health'), ('最终HP', '1420000', 'health'),
            ('证据', '', 'evidence'), ('证据', '  ', 'evidence'),
        ]
        for field, value, reason in cases:
            with self.subTest(field=field, value=value):
                invalid = {**row, field: value}
                with self.assertRaisesRegex(AssertionError, reason):
                    importer.rotation_health_overrides([invalid], self.data['modes'], rotation_modes)
        with self.assertRaisesRegex(AssertionError, 'Duplicate rotation health override'):
            importer.rotation_health_overrides([row, copy.deepcopy(row)], self.data['modes'], rotation_modes)
        with self.assertRaisesRegex(AssertionError, 'columns'):
            importer.rotation_health_overrides([{**row, '意外字段': 1}], self.data['modes'], rotation_modes)


if __name__ == '__main__':
    unittest.main()
