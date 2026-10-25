from __future__ import annotations

import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from war3_extract.workbook import MAX_CELL_UNITS, TRUNCATION_MARKER, write_workbook


class ReferenceWorkbookTests(unittest.TestCase):
    def test_numbers_preserve_precision_and_text_never_becomes_a_formula(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "reference.xlsx"
            report = write_workbook(path, [("数值", [{
                "small": 123456789012345, "large": 12345678901234567,
                "negative_large": -12345678901234567, "decimal": Decimal("117.60000000000000000001"),
                "exact_float": 117.6, "long_float": 0.12345678901234566,
                "formula": "=1+1", "plus": "+cmd", "minus": "-cmd", "at": "@SUM(A1)",
                "nested": {"base": 12345678901234567, "scale": Decimal("0.00000000000000000001")},
            }])], ["参考资料"])
            workbook = load_workbook(path, data_only=False)
            sheet = workbook["数值"]
            cells = {sheet.cell(1, index).value: sheet.cell(2, index) for index in range(1, sheet.max_column+1)}
            self.assertEqual(cells["small"].value, 123456789012345)
            self.assertEqual(cells["small"].data_type, "n")
            self.assertEqual(cells["large"].value, "12345678901234567")
            self.assertEqual(cells["negative_large"].value, "-12345678901234567")
            self.assertEqual(cells["decimal"].value, "117.60000000000000000001")
            self.assertEqual(cells["exact_float"].value, 117.6)
            self.assertEqual(cells["long_float"].value, "0.12345678901234566")
            for key in ("formula", "plus", "minus", "at", "large", "decimal"):
                self.assertEqual(cells[key].data_type, "s")
            self.assertEqual(cells["formula"].value, "=1+1")
            self.assertEqual(cells["nested"].value, '{"base":12345678901234567,"scale":"1E-20"}')
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertEqual(sheet.auto_filter.ref, "A1:K2")
            self.assertTrue(sheet["A1"].font.bold)
            self.assertEqual(report["sheets"][0]["rows"], 1)
            self.assertIn("说明", workbook.sheetnames)
            workbook.close()

    def test_sheet_names_are_valid_short_and_case_insensitive_unique(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "names.xlsx"
            long_name = "很长的表名称" * 10
            report = write_workbook(path, [("A/B", []), ("A?B", []), ("a_b", []),
                (long_name, []), (long_name, []), ("说明", [])], ["notes"])
            names = [row["name"] for row in report["sheets"]]
            self.assertEqual(names[:3], ["A_B", "A_B (2)", "a_b (3)"])
            self.assertEqual(len({name.casefold() for name in names}), len(names))
            self.assertTrue(all(len(name.encode("utf-16-le")) // 2 <= 31 for name in names))
            self.assertEqual(names[-1], "说明 (2)")
            workbook = load_workbook(path)
            self.assertEqual(workbook.sheetnames, names)
            self.assertEqual(workbook[names[0]]["A1"].value, "（无数据）")
            workbook.close()

    def test_long_cells_are_marked_and_later_row_columns_are_preserved(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "long.xlsx"
            text = "界" * 32000 + "😀" * 1000
            report = write_workbook(path, [("data", [{"long": text}, {"later": "\x00value"}])], [])
            workbook = load_workbook(path)
            sheet = workbook["data"]
            self.assertEqual([cell.value for cell in sheet[1]], ["long", "later"])
            self.assertLessEqual(len(sheet["A2"].value.encode("utf-16-le")) // 2, MAX_CELL_UNITS)
            self.assertTrue(sheet["A2"].value.endswith(TRUNCATION_MARKER))
            self.assertIsNotNone(sheet["A2"].comment)
            self.assertEqual(sheet["B3"].value, "\\u0000value")
            self.assertEqual(report["truncated_cells"], 1)
            self.assertEqual(report["total_rows"], 2)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
