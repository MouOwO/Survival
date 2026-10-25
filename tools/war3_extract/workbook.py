"""Write plain reference workbooks without formulas or numeric truncation."""
from __future__ import annotations

import json
import math
import re
import tempfile
import unicodedata
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MAX_CELL_UNITS = 32767
MAX_ROWS = 1048576
MAX_COLUMNS = 16384
TRUNCATION_MARKER = "…[单元格过长，已截断]"
INVALID_SHEET = re.compile(r"[\[\]:*?/\\\x00-\x1f]")
INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
HEADER_FILL = PatternFill("solid", fgColor="17365D")
HEADER_FONT = Font(name="Calibri", size=11, color="FFFFFF", bold=True)


def _units(text):
    return len(text.encode("utf-16-le")) // 2


def _prefix(text, maximum):
    return text.encode("utf-16-le")[:maximum * 2].decode("utf-16-le", errors="ignore")


def _text(value):
    # Preserve control characters visibly instead of silently losing content
    # or letting an XML control character make the workbook unreadable.
    return INVALID_XML.sub(lambda match: f"\\u{ord(match[0]):04x}", value)


def _value(value):
    if isinstance(value, (dict, list, tuple)):
        return _text(json.dumps(value, ensure_ascii=False, separators=(",", ":"),
            default=str, allow_nan=False))
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, bool) or value is None:
        return value
    if isinstance(value, int):
        return str(value) if len(str(abs(value))) > 15 else value
    if isinstance(value, float):
        if not math.isfinite(value):
            return str(value)
        decimal = Decimal(str(value))
        if len(decimal.as_tuple().digits) > 15 or decimal.adjusted() >= 15:
            return str(value)
        return value
    if isinstance(value, str):
        return _text(value)
    return _text(str(value))


def _sheet_name(requested, used):
    base = INVALID_SHEET.sub("_", str(requested)).strip().strip("'") or "Sheet"
    base = _prefix(base, 31).strip("'") or "Sheet"
    name, sequence = base, 1
    while name.casefold() in used:
        sequence += 1
        suffix = f" ({sequence})"
        name = _prefix(base, 31 - len(suffix)) + suffix
    used.add(name.casefold())
    return name


def _display_width(value):
    value = str(value) if value is not None else ""
    # Long descriptions should not make the worksheet hundreds of columns
    # wide; only the first line and a bounded prefix affect presentation.
    line = value.split("\n", 1)[0][:100]
    return sum(2 if unicodedata.east_asian_width(char) in ("W", "F") else 1 for char in line)


def _cell(worksheet, value, *, header=False, wrap=False):
    cell = WriteOnlyCell(worksheet)
    truncated = False
    if isinstance(value, str):
        original_length = _units(value)
        if original_length > MAX_CELL_UNITS:
            value = _prefix(value, MAX_CELL_UNITS - _units(TRUNCATION_MARKER)) + TRUNCATION_MARKER
            cell.comment = Comment(f"原文本长 {original_length} 个 UTF-16 单位；Excel 单格上限为 32767，内容已截断。", "数据导出")
            truncated = True
        cell.value = value
        # openpyxl otherwise infers '=...' as a formula. Explicit string
        # typing preserves all text, including =, +, -, @ and leading spaces.
        cell.data_type = "s"
    else:
        cell.value = value
    cell.alignment = Alignment(vertical="top", wrap_text=wrap or header)
    if header:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    return cell, truncated


def write_workbook(path: Path, sheets: list[tuple[str, list[dict]]], notes: list[str]) -> dict:
    """Write flattened dictionaries as literal Excel data and return sheet counts.

    Column order follows each key's first appearance. Notes, when provided,
    become a separate worksheet. A streaming workbook bounds additional memory
    even when a caller supplies a large wave-statistics table.
    """
    path = Path(path)
    workbook = Workbook(write_only=True)
    workbook.properties.title = "只读数据参考"
    workbook.properties.description = "提取数据的参考副本；包含原始数值与来源，不修改游戏配置。"
    requested_sheets = list(sheets)
    if notes:
        requested_sheets.append(("说明", [{"说明": note} for note in notes]))
    if not requested_sheets:
        requested_sheets.append(("说明", []))
    used, report, total_rows, total_truncated = set(), [], 0, 0
    for requested, rows in requested_sheets:
        if len(rows) + 1 > MAX_ROWS:
            raise ValueError(f"sheet {requested!r} exceeds Excel's {MAX_ROWS-1} data-row limit")
        columns = list(dict.fromkeys(key for row in rows for key in row))
        if len(columns) > MAX_COLUMNS:
            raise ValueError(f"sheet {requested!r} exceeds Excel's {MAX_COLUMNS} column limit")
        name = _sheet_name(requested, used)
        worksheet = workbook.create_sheet(name)
        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False
        worksheet.row_dimensions[1].height = 28
        headers = columns or ["（无数据）"]
        for index, key in enumerate(headers, 1):
            width = max(12, min(64, _display_width(key) + 3))
            for row in rows[:200]:
                width = max(width, min(64, _display_width(_value(row.get(key))) + 2))
            worksheet.column_dimensions[get_column_letter(index)].width = width
        header_cells = [_cell(worksheet, _text(str(key)), header=True) for key in headers]
        worksheet.append([cell for cell, _ in header_cells])
        truncated_count = sum(truncated for _, truncated in header_cells)
        for row in rows:
            cells = [_cell(worksheet, _value(row.get(key)), wrap=requested == "说明") for key in headers]
            worksheet.append([cell for cell, _ in cells])
            truncated_count += sum(truncated for _, truncated in cells)
        worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows)+1}"
        report.append({"requested_name": str(requested), "name": name, "rows": len(rows),
            "columns": len(columns), "truncated_cells": truncated_count})
        total_rows += len(rows)
        total_truncated += truncated_count
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=path.parent, prefix=".war3_reference_", suffix=".xlsx", delete=False) as temp:
        temporary = Path(temp.name)
    try:
        workbook.save(temporary)
        temporary.replace(path)
    finally:
        if temporary.exists():
            temporary.unlink()
    return {"path": str(path.resolve()), "sheets": report, "total_rows": total_rows,
        "truncated_cells": total_truncated}
