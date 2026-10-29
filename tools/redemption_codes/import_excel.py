"""Validate operator Excel and compile hashed, server-only reward definitions."""
from __future__ import annotations

import argparse
from datetime import datetime, timedelta, timezone
import hashlib
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "server"))
from archive_backend.redemption_catalog import offers, reward_spec
from archive_backend.redemption_codes import normalize


def timestamp(value):
    if value is None or value == "":
        return None
    if isinstance(value, str):
        value = datetime.fromisoformat(value.strip())
    if not isinstance(value, datetime):
        raise ValueError("redeem_date_invalid")
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone(timedelta(hours=8)))
    return value.astimezone(timezone.utc).isoformat()


def compile_workbook(path):
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=False)
    catalog, excluded = offers(ROOT)
    records, seen = [], set()
    try:
        for row in book["兑换码配置"].iter_rows(min_row=2):
            if all(c.value is None for c in row):
                continue
            if any(c.data_type == "f" for c in row):
                raise ValueError("redeem_formula_not_allowed")
            enabled, raw, label, selected, per_account, maximum, start, end = [c.value for c in row[:8]]
            if enabled not in (0, 1) or per_account != 1:
                raise ValueError("redeem_claim_limit_invalid")
            if type(maximum) is not int or not 0 <= maximum <= 1_000_000_000:
                raise ValueError("redeem_global_limit_invalid")
            if not isinstance(label, str) or not 1 <= len(label) <= 120:
                raise ValueError("redeem_label_invalid")
            digest = hashlib.sha256(normalize(raw).encode()).hexdigest()
            if digest in seen:
                raise ValueError("redeem_code_duplicated")
            seen.add(digest)
            if not isinstance(selected, str):
                raise ValueError("redeem_product_list_invalid")
            skus = list(catalog) if selected.strip() == "ALL_NON_VIP" else selected.replace("，", ",").split(",")
            skus = [s.strip() for s in skus]
            begin, expiry = timestamp(start), timestamp(end)
            if begin and expiry and datetime.fromisoformat(expiry) <= datetime.fromisoformat(begin):
                raise ValueError("redeem_date_range_invalid")
            records.append({"code_id": "r_" + digest[:32], "code_hash": digest, "label": label,
                            "enabled": enabled == 1, "max_redemptions": maximum,
                            "starts_at": begin, "expires_at": expiry,
                            "rewards": reward_spec(catalog, skus)})
    finally:
        book.close()
    if not 1 <= len(records) <= 10000:
        raise ValueError("redeem_registry_size_invalid")
    return {"version": 1, "codes": records}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    value = compile_workbook(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(value, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    print(json.dumps({"ok": True, "codes": len(value["codes"]), "output": str(args.output),
                      "plaintext_codes_in_output": False}, ensure_ascii=False))
