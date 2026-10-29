"""Create private operator workbook and server import definitions (no deployment)."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import secrets
import sys
from datetime import datetime, timezone

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "server"))
from archive_backend.redemption_codes import normalize
from archive_backend.redemption_catalog import offers, reward_spec


def create(output):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    catalog, excluded = offers(ROOT)
    book = Workbook()
    codes = book.active
    codes.title = "兑换码配置"
    codes.append(["启用", "兑换码", "名称", "商品SKU", "每账号次数", "全服总次数(0不限)", "开始时间(北京时间)", "截止时间(留空不限)", "说明"])
    definitions = []
    entries = [("all_non_vip", "全部非VIP奖品", "ALL_NON_VIP", list(catalog))]
    entries += [(sku, p["title"], sku, [sku]) for sku, p in catalog.items()]
    for key, label, selection, skus in entries:
        code = "GC-" + secrets.token_hex(12).upper()
        codes.append([1, code, label, selection, 1, 0, "", "", "永久权益与已拥有商品去重；相同码每账号一次"])
        digest = hashlib.sha256(normalize(code).encode()).hexdigest()
        definitions.append({"code_id": "r_" + digest[:32],
                            "code_hash": digest,
                            "label": label, "enabled": True, "max_redemptions": 0,
                            "starts_at": None, "expires_at": None,
                            "rewards": reward_spec(catalog, skus)})
    products = book.create_sheet("商品与奖品")
    products.append(["商品SKU", "商品名称", "类型", "原购买方式", "持有标记", "奖品说明"])
    for sku, p in catalog.items():
        products.append([sku, p["title"], p["product_type"], p.get("currency", "CNY"),
                         p["item_id"], p["description"].replace("U币", "G币")])
    skipped = book.create_sheet("排除项目")
    skipped.append(["商品SKU", "名称", "不支持原因"])
    for p in excluded:
        skipped.append([p["sku"], p["title"], p["reason"]])
    notes = book.create_sheet("使用说明")
    for row in [
        ["入口", "商城标题右侧“兑换码”；输入后点击兑换或取消。"],
        ["领取规则", "每账号每码一次；重试同一请求不会重复发奖。总次数0表示不限账号总数。"],
        ["全奖码", "兑换码配置第一行是全非VIP奖品码，含商城已实现商品、券、100G币和30天签到月卡。"],
        ["VIP排除", "不授予VIP等级、充值成长、VIP召唤权限、专属礼包和勋章。"],
        ["去重", "礼包拆成单品发放；已拥有的唯一商品跳过，其他奖品继续领取，已有权益不撤销。"],
        ["月卡", "已有有效月卡续期30天；永久月卡保留永久。"],
        ["服务器", "有效码仅登记在服务器。客户端不保存有效码表。此Excel含可领取权益的真实码，请妥善保存。"],
        ["修改导入", "使用 tools/redemption_codes/import_excel.py 校验并生成服务器导入文件；Excel修改本身不自动改服务器。"],
        ["停用/限额", "启用列改为0可停用；开始/截止时间使用北京时间，空截止表示不限时。"],
        ["修改码", "修改码会创建新码；停用旧码时需保留旧码所在行并将启用改为0。"],
        ["生成时间", datetime.now(timezone.utc).isoformat()]
    ]:
        notes.append(row)
    for number, sheet in enumerate(book.worksheets):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="334E68")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col in sheet.columns:
            letter = col[0].column_letter
            sheet.column_dimensions[letter].width = min(70, max(18, max(len(str(c.value or "")) for c in col[:20]) + 2))
        if number < 3:
            table = Table(displayName="RedemptionTable" + str(number), ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
    output.mkdir(parents=True, exist_ok=True)
    workbook = output / "孤城逐晓_兑换码配置_20261029.xlsx"
    if workbook.exists():
        raise ValueError("workbook_already_exists")
    book.save(workbook)
    registry = output / "server_codes.json"
    registry.write_text(json.dumps({"version": 1, "codes": definitions}, ensure_ascii=False,
                                   separators=(",", ":")) + "\n", encoding="utf-8")
    return {"ok": True, "workbook": str(workbook), "registry": str(registry),
            "codes": len(definitions), "products": len(catalog), "excluded": len(excluded)}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    print(json.dumps(create(parser.parse_args().output), ensure_ascii=False))
